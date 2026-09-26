"""
main.py -- Orchestrates the Power BI Report Automation.

Usage:
    python main.py          # Waits until SEND_AT time from config.py, then runs
    python main.py --now    # Runs immediately (for testing)

Flow:
    1. Wait until SEND_AT (skipped with --now)
    2. Prompt for Power BI email + password in terminal
    3. Log into Power BI (headless browser)
    4. Phase 1: Export + email for ALL AOMs
    5. Phase 2: Retry failed AOMs (if any, asks for permission)
    6. Phase 3: Retry still-failed AOMs (auto, no prompt)
    7. Print per-phase + final consolidated summary
"""
import os
import sys
import time
import logging
import getpass
from datetime import datetime, date

import openpyxl

from config import (
    STOREMASTER, EXPORTS_DIR, LOG_FILE,
    EMAIL_SUBJECT, EMAIL_BODY, SEND_AT, REPORTS
)
from powerbi import PowerBIExporter
from mailer import Mailer

# == Logging ===================================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)


# == Helpers ===================================================================

def wait_until_send_time():
    """Block until SEND_AT time is reached. Shows a simple countdown."""
    now    = datetime.now()
    hh, mm = (int(x) for x in SEND_AT.split(":"))
    target = now.replace(hour=hh, minute=mm, second=0, microsecond=0)

    if target <= now:
        log.info(f"Send time {SEND_AT} has already passed today -- running immediately.")
        return

    wait_secs = (target - now).total_seconds()
    log.info(f"Scheduled to run at {SEND_AT}. Waiting {int(wait_secs // 60)} min {int(wait_secs % 60)} sec...")
    log.info("(Run  python main.py --now  to skip the wait and test immediately.)")

    while True:
        remaining = (target - datetime.now()).total_seconds()
        if remaining <= 0:
            break
        if int(remaining) % 60 == 0:
            print(f"  ... {int(remaining // 60)} min remaining until {SEND_AT}", flush=True)
        time.sleep(1)

    log.info(f"Send time reached. Starting automation.")


def prompt_credentials():
    """Ask for Power BI email and password in the terminal."""
    print()
    print("=" * 55)
    print("  Power BI Login")
    print("=" * 55)
    email    = input("  Email    : ").strip()
    password = getpass.getpass("  Password : ")
    print("=" * 55)
    print()
    if not email or not password:
        log.error("Email or password cannot be empty.")
        sys.exit(1)
    return email, password


def load_aoms():
    """Read Store Master.xlsx and return a list of unique AOMs."""
    wb  = openpyxl.load_workbook(STOREMASTER, data_only=True)
    ws  = wb.active

    # Unhide all rows and columns unconditionally
    changed = False
    for rd in ws.row_dimensions.values():
        if rd.hidden:
            rd.hidden = False
            changed = True
    for cd in ws.column_dimensions.values():
        if cd.hidden:
            cd.hidden = False
            changed = True
    if changed:
        wb.save(STOREMASTER)
        wb = openpyxl.load_workbook(STOREMASTER, data_only=True)
        ws = wb.active

    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    idx_aom        = hdr.index("AOM")
    idx_fmail      = hdr.index("AOM Mail Id")
    idx_auto_email = hdr.index("AutoEmail")

    seen, aoms = set(), []
    for row in ws.iter_rows(min_row=2, values_only=True):
        aom_name   = row[idx_aom]
        aom_mail   = row[idx_fmail]
        auto_email = row[idx_auto_email]
        if not (aom_name and aom_mail and auto_email):
            continue
        if aom_name in seen:
            continue
        seen.add(aom_name)
        aoms.append({
            "name":           str(aom_name).strip(),
            "filter_email":   str(aom_mail).strip(),
            "filter_name":    str(aom_name).strip(),
            "delivery_email": str(auto_email).strip(),
        })
    return aoms


def cleanup_pdfs():
    """Delete all PDFs from exports/ folder."""
    if not os.path.exists(EXPORTS_DIR):
        return
    for f in os.listdir(EXPORTS_DIR):
        if f.lower().endswith(".pdf"):
            try:
                os.remove(os.path.join(EXPORTS_DIR, f))
            except OSError:
                pass


def print_phase_summary(phase_num: int, succeeded: list, failed: list):
    """Print a formatted summary block for a single phase."""
    w = 52
    log.info("")
    log.info("╔" + "═" * w + "╗")
    log.info(f"║  PHASE {phase_num} SUMMARY" + " " * (w - 16) + "║")
    log.info("╠" + "═" * w + "╣")

    if succeeded:
        names = ", ".join(succeeded)
        log.info(f"║  ✅ Sent     ({len(succeeded):>2}): {names[:w-18]}" +
                 (" ..." if len(names) > w - 18 else "") +
                 " " * max(0, w - 18 - min(len(names), w - 18)) + "  ║")
        # overflow lines
        remaining = names[w - 18:]
        while remaining:
            chunk = remaining[:w - 6]
            log.info(f"║      {chunk:<{w-6}}  ║")
            remaining = remaining[w - 6:]
    else:
        log.info(f"║  ✅ Sent     ( 0): —" + " " * (w - 22) + "  ║")

    if failed:
        names = ", ".join(failed)
        log.info(f"║  ❌ Failed   ({len(failed):>2}): {names[:w-18]}" +
                 (" ..." if len(names) > w - 18 else "") +
                 " " * max(0, w - 18 - min(len(names), w - 18)) + "  ║")
        remaining = names[w - 18:]
        while remaining:
            chunk = remaining[:w - 6]
            log.info(f"║      {chunk:<{w-6}}  ║")
            remaining = remaining[w - 6:]
    else:
        log.info(f"║  ❌ Failed   ( 0): —" + " " * (w - 22) + "  ║")

    log.info("╚" + "═" * w + "╝")
    log.info("")


def print_final_summary(total: int,
                        p1_ok: list, p1_fail: list,
                        p2_ok: list, p2_fail: list,
                        p3_ok: list, p3_fail: list):
    """Print the consolidated final summary across all phases."""
    all_sent   = p1_ok + p2_ok + p3_ok
    all_failed = p3_fail if p3_fail is not None else (p2_fail if p2_fail is not None else p1_fail)
    w = 52

    log.info("")
    log.info("╔" + "═" * w + "╗")
    log.info(f"║  FINAL SUMMARY" + " " * (w - 15) + "║")
    log.info("╠" + "═" * w + "╣")
    log.info(f"║  Total AOMs     : {total:<33}║")
    log.info(f"║  ✅ Emails sent : {len(all_sent):<33}║")
    log.info(f"║  ❌ Manual reqd : {len(all_failed):<33}║")

    if p2_ok or p2_fail is not None:
        log.info("╠" + "═" * w + "╣")
        log.info(f"║  Phase 1 sent: {len(p1_ok):<2}  Phase 1 failed: {len(p1_fail):<{w-32}}║")
        if p2_ok or p2_fail is not None:
            log.info(f"║  Phase 2 sent: {len(p2_ok):<2}  Phase 2 failed: {len(p2_fail):<{w-32}}║")
        if p3_ok or p3_fail is not None:
            log.info(f"║  Phase 3 sent: {len(p3_ok):<2}  Phase 3 failed: {len(p3_fail):<{w-32}}║")

    if all_failed:
        log.info("╠" + "═" * w + "╣")
        log.info(f"║  Please send these manually:" + " " * (w - 28) + "║")
        for name in all_failed:
            log.info(f"║    → {name:<{w-6}}║")
    log.info("╚" + "═" * w + "╝")
    log.info("")


# == Core export loop ==========================================================

def run_phase(exporter: "PowerBIExporter", mailer: "Mailer",
              aoms: list, all_aoms: list, today: str, phase_num: int) -> tuple:
    """
    Run a single export+email phase for the given list of AOMs.

    Returns (succeeded_names, failed_names) — both are plain lists of AOM names.
    """
    succeeded, failed = [], []

    for aom in aoms:
        name  = aom["name"]
        dmail = aom["delivery_email"]

        log.info(f"\nPhase {phase_num} — Processing AOM: {name}")

        pdfs = []
        for report_cfg in REPORTS:
            report_name   = report_cfg["name"]
            filter_column = report_cfg.get("filter_column", "AOM")

            filter_value = (aom["filter_email"] if filter_column == "AOM Mail Id"
                            else aom["filter_name"])

            other_filters = [
                (a["filter_email"] if filter_column == "AOM Mail Id" else a["filter_name"])
                for a in all_aoms if a["name"] != name
            ]

            log.info(f"  -> Exporting report: {report_name} (filter: {filter_column} = {filter_value})")

            pdf = exporter.export_report(
                filter_email=filter_value,
                aom_name=name,
                date_str=today,
                other_emails=other_filters,
                report_cfg=report_cfg
            )

            if pdf:
                pdfs.append(pdf)
            else:
                log.error(f"  Export failed for report '{report_name}' -- skipping this attachment for {name}")

        if not pdfs:
            log.error(f"  All exports failed -- skipping email for {name}")
            failed.append(name)
            continue

        try:
            subject = EMAIL_SUBJECT.format(aom_name=name, date=today)
            body    = EMAIL_BODY.format(aom_name=name, date=today)
            mailer.send(dmail, subject, body, attachments=pdfs)
            log.info(f"  Email sent to {dmail} with {len(pdfs)} attachment(s)")
            succeeded.append(name)
        except Exception as e:
            log.error(f"  Email failed for {name}: {e}")
            failed.append(name)

    return succeeded, failed


# == Main ======================================================================

def main():
    run_now = "--now" in sys.argv

    os.makedirs(EXPORTS_DIR, exist_ok=True)
    cleanup_pdfs()

    log.info("=" * 55)
    log.info("  Power BI Report Automation")
    log.info(f"  {'Running immediately (--now)' if run_now else f'Scheduled at {SEND_AT}'}")
    log.info("=" * 55)

    if not run_now:
        wait_until_send_time()

    pbi_email, pbi_password = prompt_credentials()

    aoms  = load_aoms()
    today = date.today().strftime("%d-%b-%Y")
    log.info(f"Loaded {len(aoms)} AOM(s) from Store Master.xlsx")
    for a in aoms:
        log.info(f"  - {a['name']} | filter: {a['filter_name']} | send to: {a['delivery_email']}")

    total  = len(aoms)
    mailer = Mailer()

    # Initialise phase result holders
    p2_ok, p2_fail = [], None   # None = phase did not run
    p3_ok, p3_fail = [], None

    with PowerBIExporter(pbi_email, pbi_password) as exporter:

        # ── Phase 1: all AOMs ────────────────────────────────────────────────
        log.info("\n" + "=" * 55)
        log.info("  PHASE 1 — Processing all AOMs")
        log.info("=" * 55)
        p1_ok, p1_fail = run_phase(exporter, mailer, aoms, aoms, today, phase_num=1)
        print_phase_summary(1, p1_ok, p1_fail)

        # ── Phase 2: retry Phase 1 failures (ask permission) ─────────────────
        if p1_fail:
            print(f"\n  {len(p1_fail)} AOM(s) failed in Phase 1.")
            answer = input("  Run Phase 2 to retry them? [Y/n]: ").strip().lower()
            if answer in ("", "y", "yes"):
                retry_aoms = [a for a in aoms if a["name"] in p1_fail]
                log.info("\n" + "=" * 55)
                log.info("  PHASE 2 — Retrying failed AOMs")
                log.info("=" * 55)
                p2_ok, p2_fail = run_phase(exporter, mailer, retry_aoms, aoms, today, phase_num=2)
                print_phase_summary(2, p2_ok, p2_fail)

                # ── Phase 3: retry Phase 2 failures (automatic) ───────────────
                if p2_fail:
                    retry_aoms = [a for a in aoms if a["name"] in p2_fail]
                    log.info("\n" + "=" * 55)
                    log.info("  PHASE 3 — Final retry (automatic)")
                    log.info("=" * 55)
                    p3_ok, p3_fail = run_phase(exporter, mailer, retry_aoms, aoms, today, phase_num=3)
                    print_phase_summary(3, p3_ok, p3_fail)
                else:
                    log.info("  Phase 2 achieved 100% success — Phase 3 not needed.")
            else:
                log.info("  Phase 2 skipped by user.")
        else:
            log.info("  Phase 1 achieved 100% success — Phase 2 and Phase 3 not needed.")

    # ── Cleanup + final summary ───────────────────────────────────────────────
    cleanup_pdfs()
    print_final_summary(total, p1_ok, p1_fail, p2_ok, p2_fail, p3_ok, p3_fail)


if __name__ == "__main__":
    main()
